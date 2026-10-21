import subprocess
import tempfile
import json
from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime

test_action_items = [
    {
        'id': 'act_live_201',
        'description': '=1+2 formula injection defense check',
        'completed': False,
        'due_at': '2026-10-15T18:00:00Z',
        'created_at': '2026-10-04T12:00:00Z',
        'updated_at': '2026-10-04T12:30:00Z',
        'conversation_id': 'conv_999',
    },
    {
        'id': 'act_live_202',
        'description': 'Deliver AI agent workflow',
        'completed': True,
        'due_at': None,
        'created_at': '2026-10-04T10:00:00Z',
        'conversation_id': 'conv_1000',
    },
]

with tempfile.TemporaryDirectory() as td:
    in_json = Path(td) / 'tasks.json'
    out_xlsx = Path(td) / 'export.xlsx'
    in_json.write_text(json.dumps(test_action_items), encoding='utf-8')

    # Execute CLI script
    cmd = ['python', 'examples/action_items_to_xlsx.py', str(in_json), '-o', str(out_xlsx), '--overwrite']
    res = subprocess.run(cmd, capture_output=True, text=True, cwd='c:/Users/Loa/Downloads/GOAL/omi_repo/sdks/python-cli')
    assert res.returncode == 0, f'CLI failed: {res.stderr}'
    print('CLI Output:', res.stdout.strip())

    # Inspect the generated Excel file
    wb = load_workbook(str(out_xlsx))
    sheet = wb.active
    assert sheet.title == 'action_items', 'Sheet title mismatch'
    assert sheet.max_row == 3, f'Expected 3 rows (1 header + 2 data), got {sheet.max_row}'

    # Check cell formula injection defense
    desc_cell = sheet['B2']
    assert desc_cell.data_type == 's', 'Description cell not string type'
    assert desc_cell.value == '=1+2 formula injection defense check', 'Description corrupted'

    # Check datetime cell
    due_cell = sheet['D2']
    assert isinstance(due_cell.value, datetime), 'Due date not native datetime'
    assert due_cell.value.hour == 18, 'Due date hour mismatch'

    wb.close()
    print('Action Items to XLSX Validation 3: PASSED 100% SUCCESS!')
